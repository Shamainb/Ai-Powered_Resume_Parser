import pytest
from docx import Document
from fpdf import FPDF
from tempfile import NamedTemporaryFile
import os

FILE_CONTENT: str = "This is a simple document."

@pytest.fixture
def generate_pdf_file():
    with NamedTemporaryFile(suffix=".pdf", mode="wb", delete=False) as temp_file:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        pdf.cell(200, 100, txt=FILE_CONTENT, align="C")
        pdf.output(temp_file.name)
        print("file name", temp_file.name)
    yield temp_file.name
    
    #os.remove(temp_file.name)
    
@pytest.fixture
def generate_doc_file():
        with NamedTemporaryFile(suffix=".docx", mode="wb", delete=False) as temp_file:
            doc = Document()
            doc.add_paragraph(FILE_CONTENT)
            doc.save(temp_file.name)
            
            temp_file.close()
            
            yield temp_file.name
            
            os.remove(temp_file.name)
